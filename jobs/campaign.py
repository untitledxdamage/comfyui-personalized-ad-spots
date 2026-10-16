#!/usr/bin/env python3
"""Run a whole customer list through the personalized-spot workflow with the Comfy SDK.

A campaign is a folder holding one small SQLite file (the ledger) with a row per person. Every
step reads and updates that ledger, so each step can be stopped, repeated, or run on another
machine:

    import     load a customer list (CSV, TSV or XLSX), clean it, reject what cannot be rendered
    submit     keep N jobs in flight until every row has been rendered
    download   fetch the finished MP4s with parallel workers, verifying their size
    run        submit + download together
    status     counts per state and an estimate of the time left
    export     manifest.csv: one line per person with file, job id and timing
    retry      put failed rows back in the queue

    pip install "comfy-sdk>=0.4"

    # local ComfyUI through the official proxy          # or a Comfy API deployment / Comfy Cloud
    pip install comfy-api-proxy && comfy-api-proxy start  # export COMFY_BASE_URL="https://<deployment>.run.comfy.app"
    export COMFY_BASE_URL="http://127.0.0.1:8189"         # export COMFY_API_KEY="comfyui-..."

    python jobs/campaign.py import --campaign campaigns/loyalty --customers jobs/sample_customers.csv
    python jobs/campaign.py run    --campaign campaigns/loyalty --batch Loyalty_Offer \
        --first-frame assets/host_first_frame.png \
        --clip 50=assets/cafe_spot.mp4 --clip 60=assets/cafe_endcard.mp4

Nothing is submitted twice: each row's idempotency key is saved before the request and its job id
right after, so a job that may already exist is looked up again instead of being sent (and billed)
again. Downloading needs only the ledger and the credentials, so it can run on a different machine
from the one that submitted, for example a cloud server next to where the videos will be hosted.
"""
import argparse, csv, hashlib, io, json, os, re, sqlite3, sys, threading, time, uuid
from concurrent.futures import ThreadPoolExecutor

# Node IDs in workflows/Personalized_Spot_API.json.
NODES = {
    "name": ("1", "value"),            # spoken in the generated greeting
    "customer_id": ("2", "value"),     # output file name
    "batch": ("3", "value"),           # output sub-folder on the server
    "greeting": ("10", "value"),       # greeting script, with {name} where the name goes
    "first_frame": ("20", "image"),    # still the greeting animates
    "seed": ("30:339", "noise_seed"),  # greeting seed (derived from customer_id -> reproducible)
}
OUTPUT_NODE = "80"                     # Save sequence -> MP4

SCHEMA = """
CREATE TABLE IF NOT EXISTS people (
    customer_id TEXT PRIMARY KEY,
    name        TEXT NOT NULL,
    seed        INTEGER NOT NULL,
    n           INTEGER NOT NULL,          -- position in the list, used for sub-folders
    status      TEXT NOT NULL DEFAULT 'pending',
    idem_key    TEXT, job_id TEXT, output_id TEXT, size_bytes INTEGER, file TEXT, error TEXT,
    attempts    INTEGER NOT NULL DEFAULT 0,
    submitted_at REAL, finished_at REAL, downloaded_at REAL
);
CREATE INDEX IF NOT EXISTS people_status ON people (status);
CREATE TABLE IF NOT EXISTS meta (key TEXT PRIMARY KEY, value TEXT);
"""
# pending -> claimed (key saved) -> submitted (job id saved) -> succeeded (rendered) -> downloaded
#                                                            \-> failed / needs_review
STATES = ["pending", "claimed", "submitted", "succeeded", "downloaded", "failed", "needs_review"]


def seed_for(customer_id: str, attempt: int = 0) -> int:
    text = customer_id if attempt == 0 else f"{customer_id}:{attempt}"
    return int.from_bytes(hashlib.sha256(text.encode("utf-8")).digest()[:6], "big")


class Ledger:
    """The campaign's state. One connection shared by all threads, guarded by a lock."""

    def __init__(self, campaign_dir, create=False):
        self.dir = campaign_dir
        path = os.path.join(campaign_dir, "campaign.db")
        if not create and not os.path.exists(path):
            raise SystemExit(f"no campaign at {campaign_dir}: run 'import' first")
        os.makedirs(campaign_dir, exist_ok=True)
        self.lock = threading.Lock()
        self.db = sqlite3.connect(path, check_same_thread=False, isolation_level=None)
        self.db.row_factory = sqlite3.Row
        self.db.execute("PRAGMA journal_mode=WAL")
        self.db.executescript(SCHEMA)

    def update(self, customer_id, **fields):
        cols = ", ".join(f"{k} = ?" for k in fields)
        with self.lock:
            self.db.execute(f"UPDATE people SET {cols} WHERE customer_id = ?", (*fields.values(), customer_id))

    def rows(self, status, limit=None):
        q = "SELECT * FROM people WHERE status = ? ORDER BY n" + (f" LIMIT {int(limit)}" if limit else "")
        with self.lock:
            return self.db.execute(q, (status,)).fetchall()

    def counts(self):
        with self.lock:
            got = dict(self.db.execute("SELECT status, COUNT(*) FROM people GROUP BY status").fetchall())
        return {s: got.get(s, 0) for s in STATES}

    def meta(self, key, value=None):
        with self.lock:
            if value is not None:
                self.db.execute("INSERT OR REPLACE INTO meta VALUES (?, ?)", (key, value))
                return value
            row = self.db.execute("SELECT value FROM meta WHERE key = ?", (key,)).fetchone()
        return row[0] if row else None


# ---------------------------------------------------------------- import
def read_table(path, delimiter=None, sheet=None):
    """Rows of a CSV / TSV / XLSX file as dicts keyed by the header line."""
    if path.lower().endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
        except ImportError:
            raise SystemExit("reading .xlsx needs openpyxl: pip install openpyxl (or export the sheet as CSV)")
        ws = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = ws[sheet] if sheet else ws.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it)]
        for values in it:
            yield {h: ("" if v is None else str(v)) for h, v in zip(header, values)}
        return
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if delimiter is None:
        first = text.split("\n", 1)[0]
        delimiter = max(",;\t|", key=first.count)
    yield from csv.DictReader(io.StringIO(text), delimiter=delimiter)


def clean_name(text, first_only, title_case):
    name = re.sub(r"\s+", " ", (text or "").strip())
    if first_only and name:
        name = name.split(" ")[0]
    if title_case:
        name = " ".join(w.capitalize() for w in name.split(" "))
    return name


def cmd_import(a):
    led = Ledger(a.campaign, create=True)
    with led.lock:
        start = led.db.execute("SELECT COALESCE(MAX(n) + 1, 0) FROM people").fetchone()[0]
        known = {r[0] for r in led.db.execute("SELECT customer_id FROM people")}
    added, rejected, seen, new = 0, [], set(known), []
    for i, row in enumerate(read_table(a.customers, a.delimiter, a.sheet), start=2):  # line 1 is the header
        if a.id_column not in row or a.name_column not in row:
            raise SystemExit(f"columns not found: need '{a.id_column}' and '{a.name_column}', file has {list(row)}")
        raw_id = (row[a.id_column] or "").strip()
        cid = re.sub(r"[^A-Za-z0-9._-]+", "_", raw_id).strip("._")     # it becomes a file name
        name = clean_name(row[a.name_column], a.first_name_only, a.title_case)
        reason = ("empty id" if not cid else "empty name" if not name else
                  f"name longer than {a.max_name_length} characters" if len(name) > a.max_name_length else
                  "name has characters that would break the script" if re.search(r'["{}\\\n]', name) else
                  "duplicate id" if cid in seen else None)
        if reason:
            if not (reason == "duplicate id" and cid in known):          # re-importing the same list is not an error
                rejected.append({"line": i, "customer_id": raw_id, "name": row[a.name_column], "reason": reason})
            continue
        seen.add(cid)
        seed = row.get(a.seed_column) if a.seed_column else None
        new.append((cid, name, int(seed) if seed else seed_for(cid), start + added))
        added += 1
    with led.lock:
        led.db.execute("BEGIN")
        led.db.executemany("INSERT INTO people (customer_id, name, seed, n) VALUES (?, ?, ?, ?)", new)
        led.db.execute("COMMIT")
    if rejected:
        path = os.path.join(a.campaign, "rejected.csv")
        with open(path, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, ["line", "customer_id", "name", "reason"])
            w.writeheader()
            w.writerows(rejected)
        print(f"{len(rejected)} rows rejected, listed with the reason in {path}")
    print(f"{added} people added, {len(known)} were already in the campaign. Total {len(known) + added}.")


# ---------------------------------------------------------------- submit
def parse_value(text):
    try:
        return json.loads(text)          # numbers, true/false
    except ValueError:
        return text


def batch_inputs(client, a, led):
    """(node, input, value) triples shared by every job. Files are uploaded once, by content hash."""
    settings = {"batch": a.batch, "greeting": a.greeting, "first_frame": a.first_frame, "clip": a.clip, "set": a.set,
                "workflow": os.path.basename(a.workflow)}
    stored = led.meta("settings")
    if stored and json.loads(stored) != settings and not a.allow_changes:
        raise SystemExit("this campaign was started with other settings:\n  " + stored +
                         "\nVideos already rendered would not match the new ones. Use the same options, "
                         "start a new campaign folder, or pass --allow-changes.")
    led.meta("settings", json.dumps(settings))
    inputs = [(*NODES["batch"], a.batch)]
    if a.greeting:
        inputs.append((*NODES["greeting"], a.greeting))
    if a.first_frame:
        inputs.append((*NODES["first_frame"], client.assets.from_file(a.first_frame)))
    for item in a.clip:
        node, path = item.split("=", 1)
        inputs.append((node, "file", client.assets.from_file(path)))
    for item in a.set:
        target, value = item.split("=", 1)
        node, field = target.rsplit(".", 1)
        inputs.append((node, field, parse_value(value)))
    return inputs


def submit_loop(client, led, a, stop=None):
    from comfy_sdk import ComfyError, IdempotencyKeyReuse, InsufficientCredits, QueueFull

    template = json.load(open(a.workflow, encoding="utf-8"))
    shared = batch_inputs(client, a, led)
    flying, budget = {}, a.limit or None
    out_of_credits, pause_until, errors_in_a_row = False, 0.0, 0

    def submit(row):
        cid = row["customer_id"]
        wf = client.workflows.from_json(json.loads(json.dumps(template)))
        for node, field, value in shared:
            wf.set_input(node, field, value)
        wf.set_input(*NODES["name"], row["name"])
        wf.set_input(*NODES["customer_id"], cid)
        wf.set_input(*NODES["seed"], row["seed"])
        key = row["idem_key"] or str(uuid.uuid4())
        led.update(cid, status="claimed", idem_key=key, error=None)       # saved BEFORE the request
        job = client.submit(wf, idempotency_key=key)
        led.update(cid, status="submitted", job_id=job.id, submitted_at=time.time(), attempts=row["attempts"] + 1)
        return job

    # Jobs a previous run left in flight: look them up, never resubmit.
    for row in led.rows("submitted"):
        try:
            flying[row["customer_id"]] = client.jobs.get(row["job_id"])
        except ComfyError as e:
            led.update(row["customer_id"], status="needs_review", error=f"job {row['job_id']} could not be looked up: {e}"[:300])
    if flying:
        print(f"resumed {len(flying)} job(s) already in flight", flush=True)

    while not (stop and stop.is_set()):
        # 1. poll what is in flight
        for cid, job in list(flying.items()):
            try:
                job.refresh()
            except Exception as e:                                         # network hiccup: try again next round
                print(f"[poll] {cid}: {str(e)[:120]}", flush=True)
                continue
            if job.status == "succeeded":
                outs = job.get_outputs(OUTPUT_NODE)
                if outs:
                    led.update(cid, status="succeeded", output_id=outs[0].id, size_bytes=outs[0].size_bytes, finished_at=time.time())
                else:
                    led.update(cid, status="failed", error=f"job succeeded but node {OUTPUT_NODE} produced no output", finished_at=time.time())
            elif job.status in ("failed", "canceled", "expired"):
                led.update(cid, status="failed", error=f"{job.status}: {job.error}"[:300], finished_at=time.time())
            else:
                continue
            del flying[cid]
            c = led.counts()
            print(f"[{job.status}] {cid}  ({c['succeeded'] + c['downloaded']} rendered, {c['pending'] + c['claimed'] + c['submitted']} to go)", flush=True)
        # 2. fill the free slots
        if not out_of_credits and time.time() >= pause_until:
            free = a.in_flight - len(flying)
            todo = (led.rows("claimed", free) + led.rows("pending", free))[:free] if free > 0 else []
            for row in todo:
                if budget is not None and budget <= 0:
                    break
                cid = row["customer_id"]
                try:
                    flying[cid] = submit(row)
                    errors_in_a_row = 0
                    if budget is not None:
                        budget -= 1
                except IdempotencyKeyReuse:
                    # An earlier attempt under this key reached the server but its job id was never saved.
                    led.update(cid, status="needs_review", error="key already used, job id unknown: check the job list before retrying")
                    print(f"[needs_review] {cid}", flush=True)
                except QueueFull:
                    led.update(cid, status="pending")                      # refused, so the key is still unused
                    pause_until = time.time() + 30
                    break
                except InsufficientCredits as e:
                    led.update(cid, status="pending")
                    out_of_credits = True
                    print(f"Out of credits ({e}). No more jobs will be submitted; the ones in flight will finish.", flush=True)
                    break
                except ComfyError as e:
                    led.update(cid, status="failed", error=str(e)[:300])
                    print(f"[error] {cid} {str(e)[:200]}", flush=True)
                    errors_in_a_row += 1
                    if errors_in_a_row >= 5:                               # a broken setup would fail every row
                        raise SystemExit("5 submissions failed in a row: stopping. Fix the cause, then 'retry' and run again.")
        c = led.counts()
        nothing_left = out_of_credits or (budget is not None and budget <= 0) or not (c["pending"] + c["claimed"])
        if not flying and nothing_left:
            break
        time.sleep(a.poll)
    return 2 if out_of_credits else 0


# ---------------------------------------------------------------- download
def target_path(out_dir, row, shard_size):
    sub = f"{row['n'] // shard_size:04d}" if shard_size else ""
    return os.path.join(out_dir, sub, f"{row['customer_id']}.mp4")


def download_one(client, led, row, out_dir, shard_size):
    """Fetch one person's MP4. Returns True when the file is on disk and verified."""
    cid = row["customer_id"]
    dst = target_path(out_dir, row, shard_size)
    try:
        job = client.jobs.get(row["job_id"])
        if row["status"] == "submitted":                                   # ledger copied while the job was running
            if job.status in ("failed", "canceled", "expired"):
                led.update(cid, status="failed", error=f"{job.status}: {job.error}"[:300], finished_at=time.time())
                return False
            if job.status != "succeeded":
                return False                                               # still rendering: next pass
            led.update(cid, status="succeeded", finished_at=time.time())
        outs = job.get_outputs(OUTPUT_NODE)
        if not outs:
            raise RuntimeError("the job no longer lists its output (retention expired?)")
        os.makedirs(os.path.dirname(dst), exist_ok=True)
        outs[0].to_file(dst + ".part")
        size = os.path.getsize(dst + ".part")
        if outs[0].size_bytes and size != outs[0].size_bytes:
            raise RuntimeError(f"incomplete download: {size} of {outs[0].size_bytes} bytes")
        os.replace(dst + ".part", dst)
        led.update(cid, status="downloaded", file=os.path.relpath(dst, out_dir).replace(os.sep, "/"), output_id=outs[0].id, size_bytes=size,
                   downloaded_at=time.time(), error=None)
        return True
    except Exception as e:
        led.update(cid, error=f"download: {e}"[:300])                      # state unchanged: a later run tries again
        print(f"[download failed] {cid}: {str(e)[:160]}", flush=True)
        return None


def download_loop(client, led, a, done=None, poll_submitted=True):
    """Download every rendered video with parallel workers.

    done: an Event set by the submit loop when it ends ("run"); None when downloading on its own.
    poll_submitted: also check rows whose job was still running when the ledger was last written.
    """
    out_dir = a.out or os.path.join("downloads", led.meta("batch") or os.path.basename(os.path.normpath(a.campaign)))
    os.makedirs(out_dir, exist_ok=True)
    led.meta("out_dir", os.path.abspath(out_dir))
    total, gave_up = 0, set()
    with ThreadPoolExecutor(a.workers) as pool:
        while True:
            last_pass = done is None or done.is_set()                      # read before the rows, so nothing slips in between
            rows = led.rows("succeeded") + (led.rows("submitted") if poll_submitted else [])
            rows = [r for r in rows if r["customer_id"] not in gave_up]
            results = list(pool.map(lambda r: download_one(client, led, r, out_dir, a.shard_size), rows))
            got = sum(1 for ok in results if ok)
            gave_up.update(r["customer_id"] for r, ok in zip(rows, results) if ok is None)
            if got:
                total += got
                print(f"downloaded {total} this run ({led.counts()['downloaded']} in the campaign) -> {out_dir}", flush=True)
            waiting = any(ok is False for ok in results)                   # jobs still rendering
            if done is not None:
                if last_pass:
                    break
            elif not (a.watch and waiting):
                break
            time.sleep(a.poll)
    return out_dir


# ---------------------------------------------------------------- commands
def client_for(led):
    from comfy_sdk import Comfy
    print(f"target: {os.environ.get('COMFY_BASE_URL') or 'Comfy Cloud'}   campaign: {led.dir}   {led.counts()}", flush=True)
    return Comfy(api_key=os.environ.get("COMFY_API_KEY"))


def cmd_submit(a):
    led = Ledger(a.campaign)
    led.meta("batch", a.batch)
    with client_for(led) as client:
        code = submit_loop(client, led, a)
    write_manifest(led)
    print(led.counts())
    return code


def cmd_download(a):
    led = Ledger(a.campaign)
    with client_for(led) as client:
        download_loop(client, led, a)
    write_manifest(led)
    print(led.counts())


def cmd_run(a):
    led = Ledger(a.campaign)
    led.meta("batch", a.batch)
    done, result = threading.Event(), {}
    with client_for(led) as client:
        t = threading.Thread(target=lambda: result.update(out=download_loop(client, led, a, done, poll_submitted=False)), daemon=True)
        t.start()
        try:
            code = submit_loop(client, led, a)
        finally:
            done.set()
            t.join()
    write_manifest(led)
    print(led.counts())
    return code


def cmd_status(a):
    led = Ledger(a.campaign)
    c = led.counts()
    total = sum(c.values())
    print(f"{total} people: " + ", ".join(f"{v} {k}" for k, v in c.items() if v))
    with led.lock:
        t = led.db.execute("SELECT MIN(submitted_at), MAX(finished_at), COUNT(*), AVG(finished_at - submitted_at) FROM people "
                           "WHERE finished_at IS NOT NULL AND status IN ('succeeded', 'downloaded')").fetchone()
    if t[2] and t[2] > 1 and t[1] > t[0]:
        rate = t[2] / (t[1] - t[0])                                        # videos per second, as observed
        left = c["pending"] + c["claimed"] + c["submitted"]
        print(f"{t[3]:.0f} s per video on average (submit to finished), {rate * 3600:.1f} videos/hour so far"
              + (f", about {left / rate / 3600:.1f} h left at this pace" if left else ""))
    if c["failed"] or c["needs_review"]:
        with led.lock:
            for r in led.db.execute("SELECT customer_id, status, error FROM people WHERE status IN ('failed', 'needs_review') LIMIT 10"):
                print(f"  {r['status']}: {r['customer_id']}  {r['error']}")


def write_manifest(led, path=None):
    path = path or os.path.join(led.dir, "manifest.csv")
    cols = ["customer_id", "name", "status", "file", "size_bytes", "seconds", "seed", "job_id", "attempts", "error"]
    with led.lock:
        rows = led.db.execute("SELECT *, ROUND(finished_at - submitted_at) AS seconds FROM people ORDER BY n").fetchall()
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(cols)
        w.writerows([r[c] for c in cols] for r in rows)
    return path


def cmd_export(a):
    led = Ledger(a.campaign)
    print("manifest:", write_manifest(led, a.out), " videos in:", led.meta("out_dir"))


def cmd_retry(a):
    led = Ledger(a.campaign)
    states = ["failed"] + (["needs_review"] if a.include_needs_review else [])
    n = 0
    for s in states:
        for r in led.rows(s):
            fields = dict(status="pending", idem_key=None, job_id=None, output_id=None, error=None)
            if a.new_seed:
                fields["seed"] = seed_for(r["customer_id"], r["attempts"])
            led.update(r["customer_id"], **fields)
            n += 1
    if a.ids:
        for cid in a.ids:                                                  # re-render specific people, e.g. a name that came out wrong
            with led.lock:
                r = led.db.execute("SELECT * FROM people WHERE customer_id = ?", (cid,)).fetchone()
            if not r:
                print(f"unknown id {cid}")
                continue
            led.update(cid, status="pending", idem_key=None, job_id=None, output_id=None, file=None, error=None,
                       seed=seed_for(cid, max(r["attempts"], 1)))
            n += 1
    print(f"{n} row(s) back in the queue")


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0], formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    here = os.path.dirname(os.path.abspath(__file__))

    def common(p):
        p.add_argument("--campaign", required=True, help="campaign folder (holds campaign.db)")
        return p

    p = common(sub.add_parser("import", help="load a customer list into the campaign"))
    p.add_argument("--customers", required=True, help="CSV, TSV or XLSX with a header line")
    p.add_argument("--id-column", default="customer_id")
    p.add_argument("--name-column", default="name")
    p.add_argument("--seed-column", help="optional column with a fixed seed per person")
    p.add_argument("--delimiter", help="CSV delimiter (default: detected from the header line)")
    p.add_argument("--sheet", help="XLSX sheet name (default: the first one)")
    p.add_argument("--first-name-only", action="store_true", help='"MARIA LOPEZ DIAZ" -> "MARIA"')
    p.add_argument("--title-case", action="store_true", help='"MARIA" -> "Maria"')
    p.add_argument("--max-name-length", type=int, default=40)
    p.set_defaults(fn=cmd_import)

    def submit_args(p):
        p.add_argument("--batch", required=True, help="campaign / batch name: it names the output folder")
        p.add_argument("--workflow", default=os.path.join(here, "..", "workflows", "Personalized_Spot_API.json"))
        p.add_argument("--first-frame", help="still image the greeting animates (uploaded once)")
        p.add_argument("--clip", action="append", default=[], metavar="NODE=FILE",
                       help="pre-made clip for a Load Video node, e.g. 50=assets/cafe_spot.mp4 (repeat per clip; uploaded once)")
        p.add_argument("--greeting", help='greeting script, e.g. "Hola, {name}. Te guardamos algo especial."')
        p.add_argument("--set", action="append", default=[], metavar="NODE.INPUT=VALUE", help="any other workflow input, e.g. 80.crf=23")
        p.add_argument("--in-flight", type=int, default=1, help="jobs kept submitted at once; match it to the number of workers")
        p.add_argument("--limit", type=int, default=0, help="submit at most this many jobs in this run (0 = all)")
        p.add_argument("--allow-changes", action="store_true", help="accept settings that differ from the ones the campaign started with")
        p.add_argument("--poll", type=float, default=3.0, help="seconds between status checks")

    def download_args(p):
        p.add_argument("--out", help="folder for the MP4s (default: downloads/<batch>)")
        p.add_argument("--workers", type=int, default=4, help="parallel downloads")
        p.add_argument("--shard-size", type=int, default=0, help="files per sub-folder (0000/, 0001/, ...); 0 = one flat folder")

    p = common(sub.add_parser("submit", help="render: keep N jobs in flight until every row is done"))
    submit_args(p)
    p.set_defaults(fn=cmd_submit)

    p = common(sub.add_parser("download", help="fetch finished videos (can run on another machine)"))
    download_args(p)
    p.add_argument("--watch", action="store_true", help="keep running and download videos as they finish")
    p.add_argument("--poll", type=float, default=10.0)
    p.set_defaults(fn=cmd_download)

    p = common(sub.add_parser("run", help="submit + download together"))
    submit_args(p)
    download_args(p)
    p.set_defaults(fn=cmd_run, watch=False)

    p = common(sub.add_parser("status", help="counts per state and time left"))
    p.set_defaults(fn=cmd_status)

    p = common(sub.add_parser("export", help="write manifest.csv"))
    p.add_argument("--out", help="manifest path (default: <campaign>/manifest.csv)")
    p.set_defaults(fn=cmd_export)

    p = common(sub.add_parser("retry", help="put failed rows back in the queue"))
    p.add_argument("--new-seed", action="store_true", help="use a different seed for the new attempt")
    p.add_argument("--include-needs-review", action="store_true", help="also retry rows whose job id was lost (may render twice)")
    p.add_argument("--ids", nargs="*", help="re-render these people with a new seed, whatever their state")
    p.set_defaults(fn=cmd_retry)

    a = ap.parse_args()
    return a.fn(a) or 0


if __name__ == "__main__":
    sys.exit(main())
